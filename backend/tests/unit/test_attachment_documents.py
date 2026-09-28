import io
import zipfile

import docx
import pytest
from PIL import Image, ImageDraw

from app.attachments.documents import (
    MAX_PDF_PAGES,
    RENDER_MAX_EDGE,
    DocumentError,
    docx_text,
    split_pdf,
)


def text_pdf(pages: list[str], *, size: tuple[int, int] = (595, 842)) -> bytes:
    """A minimal PDF with a real text layer (Helvetica), one string per page."""
    objects: list[bytes] = [b"", b""]  # 1: catalog, 2: page tree (filled in below)
    kids = []
    for text in pages:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
        content_id = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] /Contents %d 0 R "
            b"/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 "
            b"/BaseFont /Helvetica >> >> >> >>" % (size[0], size[1], content_id)
        )
        kids.append(b"%d 0 R" % len(objects))
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(kids), len(pages))
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def scanned_pdf(pages: int = 1, size: tuple[int, int] = (1240, 1754)) -> bytes:
    """Image-only pages, like a scanner produces."""
    images = []
    for i in range(pages):
        image = Image.new("RGB", size, "white")
        ImageDraw.Draw(image).text((100, 100), f"scanned page {i + 1}", fill="black")
        images.append(image)
    out = io.BytesIO()
    images[0].save(out, "PDF", save_all=True, append_images=images[1:], resolution=150)
    return out.getvalue()


def test_text_pages_give_their_text() -> None:
    pages = split_pdf(text_pdf(["The quick brown fox jumps.", "Second page has more words."]))

    assert [(p.number, p.text, p.image) for p in pages] == [
        (1, "The quick brown fox jumps.", None),
        (2, "Second page has more words.", None),
    ]


def test_scanned_pages_are_rendered_to_capped_jpegs() -> None:
    pages = split_pdf(scanned_pdf(2))

    assert [p.text for p in pages] == [None, None]
    for page in pages:
        assert page.image is not None
        with Image.open(io.BytesIO(page.image)) as image:
            assert image.format == "JPEG"
            assert max(image.size) <= RENDER_MAX_EDGE


def test_a_page_with_a_stray_page_number_counts_as_scanned() -> None:
    (page,) = split_pdf(text_pdf(["12"]))
    assert page.text is None and page.image is not None


def test_huge_page_sizes_are_rendered_small() -> None:
    # 200 x 200 inches: at 150 dpi that would be a 30000 px bitmap.
    (page,) = split_pdf(text_pdf([""], size=(14400, 14400)))
    assert page.image is not None
    with Image.open(io.BytesIO(page.image)) as image:
        assert max(image.size) <= RENDER_MAX_EDGE


def test_too_many_pages_is_refused() -> None:
    with pytest.raises(DocumentError) as info:
        split_pdf(text_pdf(["page text long enough"] * (MAX_PDF_PAGES + 1)))
    assert info.value.code == "too_many_pages"


@pytest.mark.parametrize("data", [b"%PDF-1.7\n", b"%PDF-1.4\n1 0 obj garbage", b""])
def test_broken_pdfs_are_reported(data: bytes) -> None:
    with pytest.raises(DocumentError) as info:
        split_pdf(data)
    assert info.value.code == "document_unreadable"


def word_file() -> bytes:
    document = docx.Document()
    document.add_paragraph("Homework: correct the mistakes.")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Wrong", "Right"
    table.cell(1, 0).text, table.cell(1, 1).text = "I goed", "I went"
    document.add_paragraph("Done.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_docx_keeps_paragraphs_and_tables_in_order() -> None:
    assert docx_text(word_file()) == (
        "Homework: correct the mistakes.\n\nWrong | Right\nI goed | I went\n\nDone."
    )


def test_docx_zip_bombs_are_refused() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b" " * (60 * 1024 * 1024))
    assert len(out.getvalue()) < 1024 * 1024

    with pytest.raises(DocumentError) as info:
        docx_text(out.getvalue())
    assert info.value.code == "document_unreadable"


def test_broken_docx_is_reported() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", "<not xml")
    with pytest.raises(DocumentError):
        docx_text(out.getvalue())
